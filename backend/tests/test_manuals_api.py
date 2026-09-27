import io
import json
from pathlib import Path

import pytest
from docx import Document

from app.services import enhancer, grounding, processing
from tests.claude_mock import error_client, mock_client, sse
from tests.conftest import FIRM_A, auth

SAMPLE_PDF = (Path(__file__).resolve().parents[2] / "samples" / "studio_meridian_manual.pdf").read_bytes()

# Canned model output for the sample PDF. Section 21 is TIMESHEETS, 27 is the PTO FAQ,
# and 8 is "4.2 Making a test copy of a model".
ENHANCEMENT = {
    "modules": [
        {"title": "Timesheets", "summary": "How to submit your timesheet.", "suggested_priority": "day_1", "passages": [
            {"source_section_id": 21, "heading": "Deadline", "kind": "text",
             "content": "Timesheets are due every Friday by 5:00 PM. Late timesheets hold up billing."},
        ]},
        {"title": "How to Use BIM", "summary": "Working safely with Revit models.", "suggested_priority": "day_1", "passages": [
            {"source_section_id": 8, "heading": "Test copies", "kind": "steps",
             "content": "1. Open the central model with Detach from Central.\n2. Save it to the Sandbox folder within 2 days."},
        ]},
    ],
    "issues": [
        {"type": "contradiction", "source_section_id": 21, "excerpt": "Timesheets are due every Friday by 5:00 PM",
         "related_section_id": 21, "description": "The same section later says Monday at 12:00 noon."},
        {"type": "broken_link", "source_section_id": 27, "excerpt": "go/pto-form (old link - will be updated)",
         "related_section_id": None, "description": "The manual says this link is old."},
    ],
}
GROUNDING = {"unsupported": [{"passage_id": "P2", "text": "within 2 days", "reason": "The source sets no time limit."}]}


@pytest.fixture
def claude_requests():
    return []


@pytest.fixture
def client(api, claude_requests, monkeypatch):
    claude = mock_client([sse(ENHANCEMENT), sse(GROUNDING, input_tokens=300, output_tokens=40)], claude_requests)
    monkeypatch.setattr(processing, "enhance_manual", lambda s, module_topics=None: enhancer.enhance_manual(s, module_topics=module_topics, client=claude))
    monkeypatch.setattr(processing, "check_enhancement", lambda e, s: grounding.check_enhancement(e, s, client=claude))
    return api


def upload(client, data=SAMPLE_PDF, filename="handbook.pdf", token="admin-a", **form):
    return client.post("/api/manuals", files={"file": (filename, data)}, data=form, headers=auth(token))


def docx_bytes():
    doc = Document()
    doc.add_heading("PTO", level=1)
    doc.add_paragraph("Request PTO two weeks ahead.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    ("headers", "code"),
    [
        ({}, 401),
        ({"Authorization": "Basic abc"}, 401),
        (auth("expired"), 401),
        (auth("ghost"), 403),  # valid login, no profile
        (auth("employee-a"), 403),  # not an admin
    ],
)
def test_admin_endpoints_require_admin(client, headers, code):
    assert client.get("/api/manuals", headers=headers).status_code == code


# ---------------------------------------------------------------------------
# Upload
# ---------------------------------------------------------------------------

def test_upload_parses_and_stores(client, repo, storage):
    res = upload(client, title="  Studio Handbook ")
    assert res.status_code == 201, res.text
    manual = res.json()
    assert manual["title"] == "Studio Handbook"
    assert (manual["file_type"], manual["page_count"], manual["status"]) == ("pdf", 23, "uploaded")

    path = f"{FIRM_A}/{manual['id']}.pdf"
    assert storage.files[path] == (SAMPLE_PDF, "application/pdf")
    stored = repo.select_one("manuals", {"id": manual["id"]})
    assert (stored["firm_id"], stored["uploaded_by"], stored["file_path"]) == (FIRM_A, "user-admin-a", path)

    sections = repo.select("source_sections", {"manual_id": manual["id"]}, order="ordinal")
    assert len(sections) == 58
    assert sections[8]["heading"] == "4. BIM / REVIT › 4.2 Making a test copy of a model"
    assert (sections[8]["page_start"], sections[8]["page_end"]) == (5, 5)


def test_upload_docx_uses_filename_as_title(client):
    res = upload(client, data=docx_bytes(), filename="Old Manual.docx")
    assert res.status_code == 201, res.text
    assert (res.json()["title"], res.json()["file_type"]) == ("Old Manual", "docx")


@pytest.mark.parametrize(
    ("data", "filename", "code", "message"),
    [
        (b"hello", "notes.txt", 415, "PDF or DOCX"),
        (b"PK\x03\x04 not a pdf", "manual.pdf", 415, "doesn't look like a PDF"),
        (b"%PDF-1.7 garbage", "manual.pdf", 422, "Could not open PDF"),
    ],
)
def test_upload_rejects_bad_files(client, repo, storage, data, filename, code, message):
    res = upload(client, data=data, filename=filename)
    assert res.status_code == code
    assert message in res.json()["detail"]
    assert repo.select("manuals") == [] and storage.files == {}


def test_upload_rejects_large_files(client, monkeypatch):
    from app.routers import manuals

    monkeypatch.setattr(manuals, "MAX_UPLOAD_BYTES", 1000)
    assert upload(client).status_code == 413


def test_failed_save_cleans_up(client, repo, storage):
    repo.fail_on_insert = "source_sections"
    res = upload(client)
    assert res.status_code == 500
    assert repo.select("manuals") == [] and storage.files == {}


# ---------------------------------------------------------------------------
# Listing and firm isolation
# ---------------------------------------------------------------------------

def test_list_and_get_are_scoped_to_the_firm(client):
    first = upload(client, title="First").json()
    second = upload(client, title="Second").json()

    assert [m["title"] for m in client.get("/api/manuals", headers=auth()).json()] == ["Second", "First"]
    assert client.get("/api/manuals", headers=auth("admin-b")).json() == []

    assert client.get(f"/api/manuals/{first['id']}", headers=auth()).json()["title"] == "First"
    for path in ("", "/review", "/issues"):
        assert client.get(f"/api/manuals/{second['id']}{path}", headers=auth("admin-b")).status_code == 404
    assert client.post(f"/api/manuals/{second['id']}/process", headers=auth("admin-b")).status_code == 404
    assert client.get("/api/manuals/not-a-uuid", headers=auth()).status_code == 404


# ---------------------------------------------------------------------------
# Processing, review, issues
# ---------------------------------------------------------------------------

def test_process_review_and_issues(client, repo, claude_requests):
    manual = upload(client).json()
    sections = repo.select("source_sections", {"manual_id": manual["id"]}, order="ordinal")
    section_id = {s["ordinal"]: s["id"] for s in sections}

    res = client.post(
        f"/api/manuals/{manual['id']}/process",
        json={"module_topics": ["Timesheets", " ", "How to Use BIM"]},
        headers=auth(),
    )
    assert res.status_code == 202, res.text
    assert res.json()["status"] == "processing"

    # TestClient runs background tasks before returning.
    done = client.get(f"/api/manuals/{manual['id']}", headers=auth()).json()
    assert done["status"] == "processed", done["error"]
    notes = done["processing_notes"]
    assert notes["module_topics"] == ["Timesheets", "How to Use BIM"]
    assert section_id[0] in notes["unused_section_ids"] and section_id[21] not in notes["unused_section_ids"]
    assert (notes["input_tokens"], notes["output_tokens"]) == (1500, 890)
    assert {"section_id": section_id[21], "text": "2022"} in notes["omissions"]

    enhance_prompt = json.loads(claude_requests[0].content)["messages"][0]["content"]
    assert "- Timesheets\n- How to Use BIM" in enhance_prompt

    review = client.get(f"/api/manuals/{manual['id']}/review", headers=auth()).json()
    assert len(review["sections"]) == 58
    assert [(m["title"], m["priority"], m["status"]) for m in review["modules"]] == [
        ("Timesheets", "day_1", "draft"),
        ("How to Use BIM", "day_1", "draft"),
    ]
    timesheet, bim = review["modules"][0]["passages"][0], review["modules"][1]["passages"][0]
    assert timesheet["source_section_id"] == section_id[21]
    assert timesheet["grounding_ok"] is True
    assert bim["grounding_ok"] is False
    assert [(s["text"], s["check"]) for s in bim["unsupported_spans"]] == [("2", "rule"), ("within 2 days", "model")]

    issues = client.get(f"/api/manuals/{manual['id']}/issues", headers=auth()).json()
    assert [(i["type"], i["pages"]) for i in issues] == [("contradiction", "p. 9"), ("broken_link", "pp. 11-12")]
    contradiction = issues[0]
    assert contradiction["section_heading"] == "6. TIMESHEETS"
    assert contradiction["related_section_id"] == section_id[21]
    assert contradiction["module_id"] == review["modules"][0]["id"]
    assert issues[1]["module_id"] is None  # no module uses the PTO FAQ

    assert client.get(f"/api/manuals/{manual['id']}/issues?status=resolved", headers=auth()).json() == []


def test_reprocessing_replaces_drafts(client, repo):
    manual = upload(client).json()
    for _ in range(2):
        assert client.post(f"/api/manuals/{manual['id']}/process", headers=auth()).status_code == 202
    assert len(repo.select("modules", {"manual_id": manual["id"]})) == 2
    assert len(repo.select("module_passages")) == 2
    assert len(repo.select("issues", {"manual_id": manual["id"]})) == 2


def test_process_refused_while_running_or_after_approval(client, repo):
    manual = upload(client).json()
    repo.update("manuals", manual["id"], {"status": "processing"})
    assert client.post(f"/api/manuals/{manual['id']}/process", headers=auth()).status_code == 409
    assert client.post(f"/api/manuals/{manual['id']}/process?force=true", headers=auth()).status_code == 202

    module = repo.select("modules", {"manual_id": manual["id"]})[0]
    repo.update("modules", module["id"], {"status": "approved"})
    res = client.post(f"/api/manuals/{manual['id']}/process", headers=auth())
    assert res.status_code == 409 and "already approved" in res.json()["detail"]


def test_model_failure_marks_manual_failed(client, repo, monkeypatch):
    broken = error_client(529, "Overloaded")
    monkeypatch.setattr(processing, "enhance_manual", lambda s, module_topics=None: enhancer.enhance_manual(s, client=broken))
    manual = upload(client).json()

    client.post(f"/api/manuals/{manual['id']}/process", headers=auth())

    failed = client.get(f"/api/manuals/{manual['id']}", headers=auth()).json()
    assert failed["status"] == "failed"
    assert "529" in failed["error"]
    assert repo.select("modules", {"manual_id": manual["id"]}) == []


def test_unexpected_failure_hides_details(client, repo, monkeypatch):
    def boom(*args, **kwargs):
        raise KeyError("internal detail")

    monkeypatch.setattr(processing, "enhance_manual", boom)
    manual = upload(client).json()
    client.post(f"/api/manuals/{manual['id']}/process", headers=auth())

    failed = client.get(f"/api/manuals/{manual['id']}", headers=auth()).json()
    assert failed["status"] == "failed"
    assert "internal detail" not in failed["error"]
