import io
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from app.services.parser import ParseError, parse_manual

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


@pytest.fixture(scope="module", params=["pdf", "docx"])
def sample(request):
    data = (SAMPLES / f"studio_meridian_manual.{request.param}").read_bytes()
    return request.param, parse_manual(data, request.param)


def section(parsed, heading_suffix):
    matches = [s for s in parsed.sections if s.heading and s.heading.endswith(heading_suffix)]
    assert len(matches) == 1, [s.heading for s in parsed.sections]
    return matches[0]


# ---------------------------------------------------------------------------
# Sample manual (both formats)
# ---------------------------------------------------------------------------

def test_sample_sections_follow_headings(sample):
    _, parsed = sample
    headings = [s.heading for s in parsed.sections]
    assert "1. Welcome" in headings
    assert "4. BIM / REVIT › 4.2 Making a test copy of a model" in headings
    assert "8. PAID TIME OFF (PTO) › PTO FAQ" in headings
    assert "10. HARASSMENT PREVENTION › Reporting" in headings
    assert [s.ordinal for s in parsed.sections] == list(range(len(parsed.sections)))
    assert all(s.content.strip() for s in parsed.sections)


def test_sample_formats_agree():
    pdf = parse_manual((SAMPLES / "studio_meridian_manual.pdf").read_bytes(), "pdf")
    docx = parse_manual((SAMPLES / "studio_meridian_manual.docx").read_bytes(), "docx")
    assert [s.heading for s in pdf.sections] == [s.heading for s in docx.sections]


def test_sample_keeps_links_and_paths_intact(sample):
    _, parsed = sample
    text = "\n".join(s.content for s in parsed.sections)
    assert "http://intranet.studiomeridian.example/staff/directory.aspx" in text
    assert r"\\SM-FS01\Standards\BIM\SM_BIM_Standards_v3.pdf" in text
    assert r"\\SM-FS02\Projects\<project number>\BIM\Sandbox\ and include" in text
    assert "https://studiomeridian.ethicsline.example/report (LINK TBD)" in text


def test_sample_keeps_step_numbers_as_written(sample):
    # The timesheet procedure skips step 4 on purpose; the enhancer must see the gap.
    _, parsed = sample
    content = section(parsed, "6. TIMESHEETS").content
    assert "\n3. For each project" in content
    assert "\n5. Enter your hours" in content
    assert "\n4. " not in content


def test_sample_tables_become_markdown(sample):
    _, parsed = sample
    contacts = section(parsed, "› Key contacts").content
    assert "| Topic | Contact | Title | Ext. |" in contacts
    assert "| Payroll, direct deposit, W-2 | Tom Brennan | Payroll & Accounting Manager | 120 |" in contacts
    assert "| OH-ADMIN | General office and admin tasks |" in section(parsed, "6. TIMESHEETS").content


def test_sample_bullets_are_normalized(sample):
    _, parsed = sample
    content = section(parsed, "› 4.3 Keynotes").content
    assert "\n- Do not edit SM_Keynotes.txt yourself." in content
    assert "•" not in content


def test_sample_contradictions_survive(sample):
    # The parser must keep both sides of a contradiction for the enhancer to flag.
    _, parsed = sample
    accrual = section(parsed, "› Accrual").content
    faq = section(parsed, "› PTO FAQ").content
    assert "15 days (120 hours)" in accrual
    assert "10 days (80 hours)" in faq


def test_pdf_page_numbers_and_cleanup():
    parsed = parse_manual((SAMPLES / "studio_meridian_manual.pdf").read_bytes(), "pdf")
    assert parsed.page_count == 23
    text = "\n".join(s.content for s in parsed.sections)
    assert "Employee Handbook | Page" not in text  # running footer removed
    assert "ﬀ" not in text and "ﬁ" not in text  # ligatures expanded
    assert "Staff" in text

    test_copy = section(parsed, "› 4.2 Making a test copy of a model")
    assert (test_copy.page_start, test_copy.page_end) == (5, 5)
    practices = section(parsed, "› 4.4 Model practices")
    assert (practices.page_start, practices.page_end) == (5, 6)
    assert "Revit 2025" in practices.content
    assert section(parsed, "› Training").page_start == 14


# ---------------------------------------------------------------------------
# Synthetic documents
# ---------------------------------------------------------------------------

def make_docx(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_auto_numbered_lists():
    def build(doc):
        doc.add_heading("Steps", level=1)
        for text in ("First", "Second", "Third"):
            doc.add_paragraph(text, style="List Number")
        doc.add_paragraph("Point", style="List Bullet")

    parsed = parse_manual(make_docx(build), "docx")
    assert parsed.sections[0].content == "1. First\n\n2. Second\n\n3. Third\n\n- Point"


def test_docx_nested_headings_and_page_breaks():
    def build(doc):
        doc.add_heading("Policies", level=1)
        doc.add_heading("Dress code", level=2)
        doc.add_paragraph("Business casual.")
        doc.add_page_break()
        doc.add_heading("PTO", level=1)
        doc.add_paragraph("Ask HR.")

    parsed = parse_manual(make_docx(build), "docx")
    assert [(s.heading, s.page_start) for s in parsed.sections] == [
        ("Policies › Dress code", 1),
        ("PTO", 2),
    ]
    assert parsed.page_count == 2


def test_docx_hyperlink_target_is_kept():
    def build(doc):
        doc.add_heading("Links", level=1)
        p = doc.add_paragraph("See the ")
        rel = p.part.relate_to(
            "https://example.com/pto",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "PTO form"
        run.append(t)
        link.append(run)
        p._p.append(link)

    parsed = parse_manual(make_docx(build), "docx")
    assert parsed.sections[0].content == "See the PTO form (https://example.com/pto)"


def test_text_before_first_heading_has_no_heading():
    parsed = parse_manual(make_docx(lambda d: d.add_paragraph("Intro text with no heading yet.")), "docx")
    assert parsed.sections[0].heading is None


def test_pdf_without_text_is_rejected():
    doc = pymupdf.open()
    doc.new_page()
    with pytest.raises(ParseError, match="scanned"):
        parse_manual(doc.tobytes(), "pdf")


@pytest.mark.parametrize("file_type", ["pdf", "docx"])
def test_corrupt_file_is_rejected(file_type):
    with pytest.raises(ParseError, match="Could not open"):
        parse_manual(b"not really a document", file_type)


def test_unsupported_type_is_rejected():
    with pytest.raises(ParseError, match="Unsupported"):
        parse_manual(b"", "txt")
