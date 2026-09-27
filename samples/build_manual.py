"""Render a sample source file into a DOCX and a PDF manual.

Usage:
    uv run --with python-docx --with pymupdf samples/build_manual.py            # Studio Meridian manual
    uv run --with python-docx --with pymupdf samples/build_manual.py baseline   # FIRM FLOW AEC Baseline Guide

Source format, one construct per line:
    TITLE / SUBTITLE <text>   cover lines
    PAGEBREAK                 start a new page
    # / ##                    headings
    > <text>                  note
    - <text>                  bullet
    3. <text>                 numbered step (number kept as written, gaps and all)
    | a | b |                 table row, first row of a table is the header
    Q: / A:                   FAQ lines
    anything else             paragraph
"""

import html
import re
import sys
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).parent
# name -> (source, output stem, footer label)
DOCUMENTS = {
    "firm": (HERE / "manual_source.txt", HERE / "studio_meridian_manual", "Studio Meridian Employee Handbook"),
    "baseline": (HERE / "aec_baseline_source.txt", HERE / "aec_baseline_guide", "FIRM FLOW AEC Baseline Guide"),
}

NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")


def parse(text):
    """Return a list of (kind, payload) blocks."""
    blocks = []
    table = None
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if table is None:
                table = []
                blocks.append(("table", table))
            table.append(cells)
            continue
        table = None
        if not line:
            continue
        if line == "PAGEBREAK":
            blocks.append(("pagebreak", None))
        elif line.startswith("TITLE "):
            blocks.append(("title", line[6:]))
        elif line.startswith("SUBTITLE "):
            blocks.append(("subtitle", line[9:]))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:]))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:]))
        elif line.startswith("> "):
            blocks.append(("note", line[2:]))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:]))
        elif m := NUMBERED.match(line):
            blocks.append(("step", (m.group(1), m.group(2))))
        elif line.startswith("Q: "):
            blocks.append(("q", line))
        else:
            blocks.append(("p", line))
    return blocks


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def add_page_number_footer(section, footer):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{footer}  |  Page ")
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def build_docx(blocks, path, footer):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    add_page_number_footer(doc.sections[0], footer)

    for kind, payload in blocks:
        if kind == "title":
            doc.add_heading(payload, level=0)
        elif kind == "subtitle":
            doc.add_paragraph(payload).runs[0].italic = True
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "h1":
            doc.add_heading(payload, level=1)
        elif kind == "h2":
            doc.add_heading(payload, level=2)
        elif kind == "note":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.add_run(payload).italic = True
        elif kind == "bullet":
            doc.add_paragraph(payload, style="List Bullet")
        elif kind == "step":
            num, text = payload
            p = doc.add_paragraph(f"{num}.  {text}")
            p.paragraph_format.left_indent = Inches(0.3)
        elif kind == "q":
            doc.add_paragraph().add_run(payload).bold = True
        elif kind == "table":
            t = doc.add_table(rows=0, cols=len(payload[0]), style="Table Grid")
            for i, row in enumerate(payload):
                cells = t.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
                    if i == 0:
                        cell.paragraphs[0].runs[0].bold = True
            doc.add_paragraph()
        else:
            doc.add_paragraph(payload)

    doc.save(path)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.35; }
h1 { font-size: 17pt; margin-top: 0; }
h2 { font-size: 13pt; margin-top: 12pt; }
.title { font-size: 26pt; font-weight: bold; margin-top: 120pt; }
.subtitle { font-size: 13pt; font-style: italic; }
.note { font-style: italic; margin-left: 18pt; }
.step { margin-left: 18pt; }
.q { font-weight: bold; margin-bottom: 0; }
table { border-collapse: collapse; margin-bottom: 8pt; }
td, th { border: 1px solid #555; padding: 3pt 5pt; font-size: 9.5pt; text-align: left; }
"""


def chunk_html(blocks):
    """Split blocks on page breaks and return one HTML string per chunk."""
    chunks, parts, bullets = [], [], []

    def flush_bullets():
        if bullets:
            parts.append("<ul>" + "".join(f"<li>{b}</li>" for b in bullets) + "</ul>")
            bullets.clear()

    for kind, payload in blocks:
        if kind != "bullet":
            flush_bullets()
        if kind == "pagebreak":
            chunks.append("".join(parts))
            parts = []
            continue
        if kind == "table":
            head, *rows = payload
            cells = "".join(f"<th>{html.escape(c)}</th>" for c in head)
            body = "".join(
                "<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in rows
            )
            parts.append(f"<table><tr>{cells}</tr>{body}</table>")
            continue
        if kind == "step":
            num, text = payload
            parts.append(f'<p class="step">{num}.&nbsp;&nbsp;{html.escape(text)}</p>')
            continue
        text = html.escape(payload)
        if kind == "bullet":
            bullets.append(text)
        elif kind in ("h1", "h2"):
            parts.append(f"<{kind}>{text}</{kind}>")
        elif kind in ("title", "subtitle", "note", "q"):
            parts.append(f'<p class="{kind}">{text}</p>')
        else:
            parts.append(f"<p>{text}</p>")

    flush_bullets()
    chunks.append("".join(parts))
    return [c for c in chunks if c]


def build_pdf(blocks, path, footer):
    page = fitz.paper_rect("letter")
    content = page + (54, 54, -54, -72)
    writer = fitz.DocumentWriter(str(path))
    for chunk in chunk_html(blocks):
        story = fitz.Story(html=chunk, user_css=CSS)
        more = True
        while more:
            device = writer.begin_page(page)
            more, _ = story.place(content)
            story.draw(device)
            writer.end_page()
    writer.close()

    doc = fitz.open(path)
    for i, pg in enumerate(doc, start=1):
        label = f"{footer}  |  Page {i}"
        pg.insert_text((page.width / 2 - 110, page.height - 36), label, fontsize=8)
    doc.saveIncr()


def main():
    source, out_stem, footer = DOCUMENTS[sys.argv[1] if len(sys.argv) > 1 else "firm"]
    blocks = parse(source.read_text())
    build_docx(blocks, out_stem.with_suffix(".docx"), footer)
    build_pdf(blocks, out_stem.with_suffix(".pdf"), footer)
    pages = fitz.open(out_stem.with_suffix(".pdf")).page_count
    print(f"Wrote {out_stem.name}.docx and {out_stem.name}.pdf ({pages} pages)")


if __name__ == "__main__":
    main()
